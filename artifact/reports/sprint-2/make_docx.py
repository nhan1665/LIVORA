"""Build the English Sprint 2 report with genuine browser image evidence."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

artifact = Path(__file__).resolve().parents[2]
source = artifact / 'SPRINT_2_IMPLEMENTATION_REPORT.md'
output = artifact / 'SPRINT_2_IMPLEMENTATION_REPORT.docx'
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.72)
section.right_margin = Inches(.72)
normal = doc.styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(9)
normal.paragraph_format.space_after = Pt(5)
for name, size in [('Title', 23), ('Heading 1', 16), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Aptos Display'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(89, 60, 59)
section.header.paragraphs[0].text = 'LIVORA  |  SPRINT 2 CUSTOMER & SETTINGS INTEGRATION'
section.footer.paragraphs[0].text = 'Implementation and verified local QA  •  2026-10-25'
section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

def plain(value):
    value = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', value)
    return value.replace('**', '').replace('`', '').strip()

lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip() or line.strip() == '---':
        i += 1
        continue
    image = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', line)
    if image:
        path = artifact / image.group(2)
        if not path.is_file():
            raise FileNotFoundError(path)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(6.55))
        caption = doc.add_paragraph(image.group(1))
        caption.style = 'Caption'
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
        continue
    heading = re.match(r'^(#{1,3})\s+(.+)$', line)
    if heading:
        doc.add_heading(plain(heading.group(2)), 0 if len(heading.group(1)) == 1 else len(heading.group(1)) - 1)
        i += 1
        continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [plain(c) for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{3,}:?', c.replace(' ', '')) for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            table = doc.add_table(rows=0, cols=max(map(len, rows)))
            table.style = 'Light Shading Accent 1'
            for row in rows:
                cells = table.add_row().cells
                for j, value in enumerate(row):
                    cells[j].text = value
            doc.add_paragraph()
        continue
    bullet = re.match(r'^[-*]\s+(.+)$', line)
    numbered = re.match(r'^\d+\.\s+(.+)$', line)
    if bullet or numbered:
        doc.add_paragraph(plain((bullet or numbered).group(1)), style='List Bullet' if bullet else 'List Number')
        i += 1
        continue
    paragraph = [line.rstrip().rstrip('\\')]
    i += 1
    while i < len(lines) and lines[i].strip() and not lines[i].startswith(('#', '|', '![', '-', '*')) and not re.match(r'^\d+\.\s+', lines[i]):
        paragraph.append(lines[i].strip())
        i += 1
    doc.add_paragraph(plain(' '.join(paragraph)))

doc.save(output)
check = Document(output)
assert len(check.inline_shapes) == 7
print(f'Generated {output.name}: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables, {len(check.inline_shapes)} real screenshots.')
