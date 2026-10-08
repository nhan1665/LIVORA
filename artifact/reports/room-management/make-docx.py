from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
artifact=Path(r'E:/web2/LIVORA/artifact')
source=artifact/'ROOM_MANAGEMENT_IMPLEMENTATION_REPORT.md'
text=source.read_text(encoding='utf-8')
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
doc.styles['Normal'].font.name='Aptos'; doc.styles['Normal'].font.size=Pt(9); doc.styles['Normal'].paragraph_format.space_after=Pt(4)
for name,size,color in [('Title',24,'593C3B'),('Heading 1',17,'593C3B'),('Heading 2',13,'704B4A'),('Heading 3',10,'704B4A')]:
    st=doc.styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
sec.header.paragraphs[0].text='LIVORA  |  COMPLETE ADMIN ROOM MANAGEMENT MODULE'; sec.header.paragraphs[0].style='Caption'
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('Implementation, QA & Figma comparison  •  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
def plain(s):
    s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',s)
    return s.replace('**','').replace('`','').replace('~~','').strip()
def shade(cell,color):
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); cell._tc.get_or_add_tcPr().append(shd)
lines=text.splitlines(); i=0; code=False; buf=[]
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        if not code: code=True; buf=[]
        else:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.space_after=Pt(7)
            run=p.add_run('\n'.join(buf)); run.font.name='Consolas'; run.font.size=Pt(8); code=False
        i+=1; continue
    if code: buf.append(line); i+=1; continue
    h=re.match(r'^(#{1,6})\s+(.*)$',line)
    if h:
        if h.group(2).startswith('19.9 Final acceptance checklist'):
            doc.add_page_break()
        level=len(h.group(1)); doc.add_heading(plain(h.group(2)),0 if level==1 else min(level-1,3)); i+=1; continue
    im=re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$',line)
    if im:
        path=(artifact/im.group(2)).resolve()
        if path.exists():
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(6.1))
            c=doc.add_paragraph(plain(im.group(1))); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.style='Caption'
        else: doc.add_paragraph('Missing screenshot: '+im.group(2))
        i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{3,}:?',c.replace(' ','')) for c in row): rows.append(row)
            i+=1
        if rows:
            cols=max(len(x) for x in rows); table=doc.add_table(rows=0,cols=cols); table.style='Light Shading Accent 1'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
            for ri,row in enumerate(rows):
                cells=table.add_row().cells
                for ci in range(cols):
                    cells[ci].text=plain(row[ci]) if ci<len(row) else ''; cells[ci].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for p in cells[ci].paragraphs:
                        p.paragraph_format.space_after=Pt(2)
                        for run in p.runs: run.font.size=Pt(8)
                    if ri==0:
                        shade(cells[ci],'704B4A')
                        for p in cells[ci].paragraphs:
                            for run in p.runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
        continue
    if re.match(r'^\s*[-*]\s+',line): doc.add_paragraph(plain(re.sub(r'^\s*[-*]\s+','',line)),style='List Bullet'); i+=1; continue
    if re.match(r'^\s*\d+\.\s+',line): doc.add_paragraph(plain(re.sub(r'^\s*\d+\.\s+','',line)),style='List Number'); i+=1; continue
    if not line.strip() or line.strip()=='---': i+=1; continue
    para=[line.strip()]; i+=1
    while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','```','![')) and not re.match(r'^\s*(?:[-*]\s+|\d+\.\s+)',lines[i]):
        para.append(lines[i].strip()); i+=1
    doc.add_paragraph(plain(' '.join(para)))
out=artifact/'ROOM_MANAGEMENT_IMPLEMENTATION_REPORT.docx'; doc.save(out)
check=Document(out)
print(f'Generated {out} ({out.stat().st_size} bytes); {len(check.paragraphs)} paragraphs, {len(check.tables)} tables, {len(check.inline_shapes)} embedded images')
